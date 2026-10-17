"""Teklif raporları — yönetici için özet, müşteri/model/şehir kırılımları."""

import io
from collections import defaultdict
from datetime import datetime, timedelta

from fastapi import APIRouter, Request
from fastapi.responses import Response

import auth
import db.factory as fdb
import db.users as udb
from tmpl import templates

router = APIRouter(prefix="/raporlar")

# Teklifin "kazanıldı" sayıldığı durumlar
WON = {"Onaylandı", "Sipariş Verildi", "Admin Onaylı", "Üretimde",
       "Tamamlandı", "Teslim Edildi"}
LOST = {"Reddedildi", "İptal"}

PRESETS = [("30", "Son 30 gün"), ("90", "Son 90 gün"), ("180", "Son 6 ay"),
           ("365", "Son 1 yıl"), ("0", "Tümü")]


def _d(s):
    """'2026-10-17 12:00:00' → date, bozuksa None."""
    if not s:
        return None
    try:
        return datetime.strptime(str(s)[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


def _money(v):
    try:
        return f"{float(v or 0):,.2f}".replace(",", "\x00").replace(".", ",").replace("\x00", ".")
    except (TypeError, ValueError):
        return "0,00"


class Bag:
    """Para birimi karışmasın diye tutarları birimine göre ayrı toplar."""

    def __init__(self):
        self.n = 0
        self.by_cur = defaultdict(float)

    def add(self, offer):
        self.n += 1
        self.by_cur[offer.get("currency") or "EUR"] += float(
            offer.get("final_price") or offer.get("total_price") or 0)

    def text(self):
        if not self.by_cur:
            return "—"
        return " + ".join(f"{_money(v)} {k}"
                          for k, v in sorted(self.by_cur.items(), key=lambda x: -x[1]))

    @property
    def main(self):
        """Sıralama için: en büyük tutarlı para biriminin değeri."""
        return max(self.by_cur.values()) if self.by_cur else 0.0


def _collect(request, days, status, city, dealer_scope):
    offers = fdb.get_offers(dealer_id=dealer_scope)
    customers = {c["id"]: c for c in fdb.get_customers()}
    models = {m["id"]: m for m in fdb.get_models()}

    since = None
    if days:
        since = (datetime.now() - timedelta(days=days)).date()

    rows = []
    for o in offers:
        cu = customers.get(o.get("customer_id")) or {}
        od = _d(o.get("created_at"))
        if since and (od is None or od < since):
            continue
        if status and o.get("status") != status:
            continue
        if city and (cu.get("city") or "") != city:
            continue
        o["customer_name"] = cu.get("name") or o.get("model_name") or "—"
        o["customer_city"] = cu.get("city") or ""
        o["customer_obj"] = cu
        o["model_disp"] = (models.get(o.get("model_id")) or {}).get("name") \
            or o.get("model_name") or "—"
        o["date"] = od
        o["amount"] = float(o.get("final_price") or o.get("total_price") or 0)
        rows.append(o)
    return rows, customers


@router.get("")
async def reports(request: Request, days: int = 90, status: str = "",
                  city: str = "", sort: str = "count"):
    user = auth.require_admin(request)
    rows, customers = _collect(request, days, status, city, None)

    # ── Özet ──
    total, won, lost, open_ = Bag(), Bag(), Bag(), Bag()
    for o in rows:
        total.add(o)
        (won if o["status"] in WON else lost if o["status"] in LOST else open_).add(o)
    conv = (won.n / total.n * 100) if total.n else 0.0

    # ── Durum dağılımı ──
    by_status = defaultdict(Bag)
    for o in rows:
        by_status[o["status"]].add(o)
    status_rows = sorted(by_status.items(), key=lambda x: -x[1].n)

    # ── Müşteri kırılımı ──
    by_cust = defaultdict(lambda: {"bag": Bag(), "offers": [], "cust": {}})
    for o in rows:
        e = by_cust[o.get("customer_id")]
        e["bag"].add(o)
        e["offers"].append(o)
        e["cust"] = o["customer_obj"]
    cust_rows = []
    for cid, e in by_cust.items():
        offs = sorted(e["offers"], key=lambda x: (x["date"] or datetime.min.date()), reverse=True)
        cust_rows.append({
            "id": cid,
            "name": e["cust"].get("name") or "—",
            "city": e["cust"].get("city") or "—",
            "contact": e["cust"].get("contact_person") or "",
            "n": e["bag"].n,
            "total": e["bag"].text(),
            "sortv": e["bag"].main,
            "won": sum(1 for x in offs if x["status"] in WON),
            "last": offs[0]["date"] if offs else None,
            "offers": offs,
            "models": sorted({x["model_disp"] for x in offs}),
        })
    keyf = {"count": lambda r: (-r["n"], -r["sortv"]),
            "amount": lambda r: -r["sortv"],
            "name": lambda r: r["name"].lower()}.get(sort, lambda r: (-r["n"], -r["sortv"]))
    cust_rows.sort(key=keyf)
    repeat_rows = [r for r in cust_rows if r["n"] >= 2]

    # ── Model kırılımı ──
    by_model = defaultdict(Bag)
    for o in rows:
        by_model[o["model_disp"]].add(o)
    model_rows = sorted(
        [{"name": k, "n": v.n, "total": v.text(), "sortv": v.main}
         for k, v in by_model.items()], key=lambda r: -r["n"])

    # ── Şehir kırılımı ──
    by_city = defaultdict(Bag)
    for o in rows:
        by_city[o["customer_city"] or "—"].add(o)
    city_rows = sorted(
        [{"name": k, "n": v.n, "total": v.text(), "sortv": v.main}
         for k, v in by_city.items()], key=lambda r: -r["n"])

    # ── Aylık seyir ──
    by_month = defaultdict(Bag)
    for o in rows:
        if o["date"]:
            by_month[o["date"].strftime("%Y-%m")].add(o)
    months = sorted(by_month.items())[-12:]
    mx = max((b.n for _, b in months), default=1) or 1
    month_rows = [{"label": datetime.strptime(k, "%Y-%m").strftime("%m/%Y"),
                   "n": b.n, "total": b.text(), "pct": round(b.n / mx * 100)}
                  for k, b in months]

    all_cities = sorted({(c.get("city") or "").strip()
                         for c in customers.values() if (c.get("city") or "").strip()})

    return templates.TemplateResponse(request, "reports.html", {
        "user": user,
        "rows": rows,
        "total": total, "won": won, "lost": lost, "open": open_,
        "conv": conv,
        "status_rows": status_rows,
        "cust_rows": cust_rows,
        "repeat_rows": repeat_rows,
        "model_rows": model_rows,
        "city_rows": city_rows,
        "month_rows": month_rows,
        "all_cities": all_cities,
        "presets": PRESETS,
        "days": days, "status": status, "city": city, "sort": sort,
        "statuses": sorted({o["status"] for o in fdb.get_offers()}),
        "active_page": "reports",
    })


@router.get("/excel")
async def reports_excel(request: Request, days: int = 90, status: str = "",
                        city: str = ""):
    """Filtrelenmiş teklif listesini Excel olarak indirir."""
    auth.require_admin(request)
    rows, _ = _collect(request, days, status, city, None)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        return Response("openpyxl kurulu değil.", status_code=500,
                        media_type="text/plain")

    wb = Workbook()
    ws = wb.active
    ws.title = "Teklifler"
    head = ["Teklif No", "Tarih", "Müşteri", "Şehir", "İlgili Kişi", "Model",
            "Adet", "Tutar", "Para Birimi", "Durum"]
    ws.append(head)
    hf = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1E2B4A")
    for i in range(1, len(head) + 1):
        cell = ws.cell(row=1, column=i)
        cell.font = hf
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")

    for o in sorted(rows, key=lambda x: (x["date"] or datetime.min.date()), reverse=True):
        cu = o["customer_obj"]
        ws.append([o.get("offer_no") or "", o["date"].strftime("%d.%m.%Y") if o["date"] else "",
                   o["customer_name"], o["customer_city"] or "",
                   cu.get("contact_person") or "", o["model_disp"],
                   int(o.get("machine_count") or 1), round(o["amount"], 2),
                   o.get("currency") or "", o.get("status") or ""])

    for col, w in zip("ABCDEFGHIJ", (20, 12, 34, 14, 20, 28, 7, 14, 12, 16)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    fn = f"teklif_raporu_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return Response(
        buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fn}"'})
